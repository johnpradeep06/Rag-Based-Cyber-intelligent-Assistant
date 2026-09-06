import os
import re
import csv
import time
import json
import shutil
import tempfile
import bs4
import requests
from dotenv import load_dotenv
from exa_py import Exa
from openai import OpenAI


from langchain_core.documents import Document
from langchain_core.prompts import PromptTemplate
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import (
    WebBaseLoader, PyPDFLoader, TextLoader, Docx2txtLoader, GitLoader,
)
from langchain_community.vectorstores import Chroma
from langchain_core.output_parsers import StrOutputParser
from langchain_core.runnables import RunnablePassthrough
from langchain_openai import ChatOpenAI, OpenAIEmbeddings
from google import genai
from google.genai import types

# =========================================================
# LOAD ENV
# =========================================================

load_dotenv()
exa = Exa(api_key=os.environ.get("EXA_API_KEY"))
# =========================================================
# LANGSMITH CONFIG
# =========================================================

os.environ["LANGCHAIN_TRACING_V2"] = "true"
os.environ["LANGCHAIN_ENDPOINT"] = "https://api.smith.langchain.com"

# =========================================================
# CONFIG
# =========================================================

RELEVANCE_THRESHOLD = 0.15
OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY")
# Same DATA_DIR convention as database.py — a persistent volume in production.
CHROMA_PERSIST_DIR = os.path.join(os.getenv("DATA_DIR", "."), "chroma_db")

if not OPENROUTER_API_KEY:
    raise ValueError("OPENROUTER_API_KEY not found in .env file")

# =========================================================
# INDEXING & STORAGE
# =========================================================

embedding_func = OpenAIEmbeddings(
    api_key=OPENROUTER_API_KEY,
    base_url="https://openrouter.ai/api/v1",
    model="openai/text-embedding-ada-002",
)

vectorstore = Chroma(
    persist_directory=CHROMA_PERSIST_DIR,
    embedding_function=embedding_func
)

text_splitter = RecursiveCharacterTextSplitter(
    chunk_size=1000,
    chunk_overlap=200,
)

def _invalidate_bm25():
    """Tell the opt-in retrieval module its lexical index is stale. No-op unless
    retrieval.py is importable (RETRIEVAL_V2 workflow)."""
    try:
        import retrieval
        retrieval.reset_bm25()
    except Exception:
        pass


def _tag_and_store(docs, *, source_type: str, title: str, origin: str, ref: str,
                   split: bool = True) -> int:
    """Stamp normalized metadata on every doc, chunk (prose only), and index.
    Returns the number of chunks written."""
    clean = []
    for d in docs:
        text = (d.page_content or "").strip()
        if not text:
            continue
        md = d.metadata or {}
        d.metadata = {
            **md,
            "source_type": source_type,
            "title": md.get("title") or title,   # per-doc title wins (e.g. a CVE id)
            "origin": origin,
            "ref": ref,
        }
        clean.append(d)
    if not clean:
        return 0
    chunks = text_splitter.split_documents(clean) if split else clean
    vectorstore.add_documents(documents=chunks)
    _invalidate_bm25()
    return len(chunks)


def _load_file_docs(path: str):
    """Return (list[Document], source_type, split?) for a local file, by extension."""
    ext = path.rsplit(".", 1)[-1].lower() if "." in path else ""
    if ext == "pdf":
        return PyPDFLoader(path).load(), "pdf", True
    if ext == "docx":
        return Docx2txtLoader(path).load(), "docx", True
    if ext == "csv":
        return _csv_docs(path), "csv", False
    if ext in ("xlsx", "xlsm"):
        return _xlsx_docs(path), "xlsx", False
    if ext == "json":
        with open(path, "r", encoding="utf-8", errors="replace") as fh:
            return _json_docs(fh.read()), "json", False
    if ext == "log":
        return _log_docs(path), "log", True
    # txt / md / rst / yaml / yml / anything else -> plain text
    st = "text"
    try:
        docs = TextLoader(path, autodetect_encoding=True).load()
    except Exception:
        with open(path, "r", encoding="utf-8", errors="replace") as fh:
            docs = [Document(page_content=fh.read(), metadata={"source": path})]
    return docs, st, True


def _csv_docs(path: str):
    out = []
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
        for i, row in enumerate(csv.DictReader(fh)):
            body = "\n".join(f"{k}: {v}" for k, v in row.items() if v not in (None, ""))
            if body:
                out.append(Document(page_content=body, metadata={"row": i + 1}))
    return out


def _xlsx_docs(path: str):
    import openpyxl
    out = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in wb.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        header = [str(h) if h is not None else f"col{j}" for j, h in enumerate(header)]
        for i, row in enumerate(rows):
            body = "\n".join(
                f"{header[j]}: {v}" for j, v in enumerate(row)
                if j < len(header) and v not in (None, "")
            )
            if body:
                out.append(Document(page_content=body,
                                    metadata={"sheet": sheet.title, "row": i + 2}))
    wb.close()
    return out


def _json_docs(text: str, cap: int = 2000):
    data = json.loads(text)
    items = data if isinstance(data, list) else [data]
    out = []
    for i, item in enumerate(items[:cap]):
        body = json.dumps(item, indent=2, ensure_ascii=False) if not isinstance(item, str) else item
        out.append(Document(page_content=body, metadata={"index": i}))
    return out


def _log_docs(path: str, max_bytes: int = 2 * 1024 * 1024):
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        raw = fh.read(max_bytes)
    # collapse runs of identical lines: "<line>  (xN)"
    lines, prev, count, kept = raw.splitlines(), None, 0, []
    for ln in lines:
        if ln == prev:
            count += 1
            continue
        if prev is not None:
            kept.append(prev if count == 1 else f"{prev}  (x{count})")
        prev, count = ln, 1
    if prev is not None:
        kept.append(prev if count == 1 else f"{prev}  (x{count})")
    return [Document(page_content="\n".join(kept),
                     metadata={"source": path, "line_range": f"1-{len(lines)}"})]


def ingest_document(file_path: str) -> int:
    """Back-compat synchronous file ingest (used by the legacy /upload endpoint)."""
    docs, st, split = _load_file_docs(file_path)
    return _tag_and_store(docs, source_type=st, title=os.path.basename(file_path),
                          origin="file", ref=os.path.basename(file_path), split=split)

# =========================================================
# RETRIEVAL
# =========================================================

def retrieve_context(question: str) -> str | None:
    results = vectorstore.similarity_search_with_relevance_scores(
        question,
        k=4,
    )

    relevant_docs = [
        doc for doc, score in results if score >= RELEVANCE_THRESHOLD
    ]

    if not relevant_docs:
        return None

    return "\n\n".join(doc.page_content for doc in relevant_docs)

# =========================================================
# GREETING DETECTOR
# =========================================================

def is_greeting(text: str) -> bool:
    greetings = {
        "hi",
        "hello",
        "hey",
        "hai",
        "hii",
        "good morning",
        "good afternoon",
        "good evening",
        "whats up",
        "what's up",
    }

    text = text.lower().strip()
    return any(text == g or text.startswith(g) for g in greetings)

# =========================================================
# PROMPT
# =========================================================

prompt = PromptTemplate(
    input_variables=["context", "question"],
    template="""
You are Sentinel, a cyber security intelligence assistant. You answer from the
retrieved context below — security reports, advisories, playbooks, detection
rules and reference material an analyst has indexed.

How to answer:
- Synthesise an answer from the context. Combine details across passages — you do
  NOT need a single passage that states the answer word for word. If the passages
  partially cover the question, give the partial answer and note what is missing.
- Write like a security analyst: precise and concise, no filler. Use short bullet
  points or numbered steps where they help.
- Ground every specific claim in the context. Do not introduce CVE numbers,
  versions, commands, function names or IOCs that are not in the context.
- Only if the context genuinely says nothing relevant to the question, reply with
  one sentence: "The indexed sources don't cover this." and stop.
- If the question is about your identity or capabilities, answer directly.

Context:
{context}

Question:
{question}

Answer:
"""
)

# =========================================================
# LLM
# =========================================================

llm = ChatOpenAI(
    api_key=OPENROUTER_API_KEY,
    base_url="https://openrouter.ai/api/v1",
    model="openai/gpt-oss-120b",
    max_tokens=1000
)

# =========================================================
# SECONDARY RAG / FALLBACK
# =========================================================

def is_university_relevant(question: str) -> bool:
    """Kept name for call-site stability; now gates on cyber-security relevance."""
    relevance_prompt = PromptTemplate(
        input_variables=["question"],
        template="""
Determine if the following question is about cyber security, information security,
threat intelligence, hacking techniques, malware, or digital defence.
Respond with exactly YES or NO.

Question: {question}
"""
    )
    chain = relevance_prompt | llm | StrOutputParser()
    try:
        result = chain.invoke({"question": question})
        return "YES" in result.strip().upper()
    except Exception:
        return False

def exa_search_fallback(question: str) -> str:
    import re
    try:
        response = exa.answer(question)
        answer_text = response.answer
        
        # Extract markdown links from Exa's response
        links = re.findall(r'\[([^\]]+)\]\((https?://[^\)]+)\)', answer_text)
        
        sources = []
        seen_urls = set()
        for title, url in links:
            if url not in seen_urls:
                sources.append(f"- [{title}]({url})")
                seen_urls.add(url)
                
        # Remove the citation blocks e.g., ([Title](url), [Title](url))
        answer_text = re.sub(r'\s*\((?:\[[^\]]+\]\((?:https?://[^\)]+)\)(?:,\s*)?)+\)', '', answer_text)
        # Also remove any remaining bare inline links like [Title](url) -> Title
        answer_text = re.sub(r'\[([^\]]+)\]\((https?://[^\)]+)\)', r'\1', answer_text)
        
        if sources:
            sources_text = "\n\n**Sources:**\n" + "\n".join(sources)
            answer_text += sources_text
            
        return answer_text
    except Exception as e:
        # Web search is a best-effort extra (needs a valid EXA_API_KEY). Never
        # surface a raw transport/401 error as the answer.
        print(f"[exa_search_fallback] disabled: {e}")
        return "I couldn't find this in the indexed sources. Add a relevant document from the admin panel and ask again."

# =========================================================
# RAG FUNCTION
# =========================================================

FALLBACK_TRIGGERS = [
    "don't know based on the given context",
    "don’t know based on the given context",
    "do not know based on the given context",
    "don't know based on the context",
    "don’t know based on the context",
    "indexed sources don't cover this",
    "indexed sources don’t cover this",
    "sources don't cover this",
    "don't know",  # catch shorter versions of the LLM defying prompt rules
]


def _is_refusal(text: str) -> bool:
    """True when the whole reply is a 'not covered' notice (not a real answer)."""
    t = (text or "").strip().lower()
    return len(t) < 120 and any(trigger in t for trigger in FALLBACK_TRIGGERS)


# =========================================================
# SHORT-TERM CONVERSATION MEMORY
# =========================================================
# The caller (app.py) passes `history` = recent [{role, content}] for the session,
# oldest first, already windowed. Two uses:
#   1. condense a follow-up into a standalone question for retrieval
#   2. give the generator a short transcript so the answer stays coherent
# Both degrade to "no memory" on any error.

MEMORY_WINDOW = int(os.getenv("MEMORY_WINDOW", "6"))


def _memory_on() -> bool:
    return os.getenv("CONVERSATION_MEMORY_ENABLED", "true").strip().lower() in (
        "1", "true", "yes", "on")


def _history_block(history) -> str:
    if not history:
        return ""
    lines = []
    for m in history:
        who = "User" if m.get("role") == "user" else "Assistant"
        c = " ".join((m.get("content") or "").split())[:500]
        if c:
            lines.append(f"{who}: {c}")
    return "\n".join(lines)


def _condense_question(history, question: str) -> str:
    """Rewrite a follow-up into a standalone question using recent turns.
    Returns `question` unchanged when there is no history or on any error."""
    block = _history_block(history)
    if not block:
        return question
    try:
        r = _stream_client.chat.completions.create(
            model="openai/gpt-oss-120b", max_tokens=160, temperature=0.0,
            messages=[{"role": "user", "content":
                       "Given the conversation and a follow-up question, rewrite the "
                       "follow-up as a standalone question that makes sense on its own. "
                       "Keep it faithful and do NOT answer it. If it is already "
                       "standalone, return it unchanged.\n\n"
                       f"Conversation:\n{block}\n\nFollow-up: {question}\n\n"
                       "Standalone question:"}],
            # gpt-oss-120b always reasons; `exclude` keeps that out of `content`.
            extra_body={"reasoning": {"effort": "low", "exclude": True}},
        )
        return (r.choices[0].message.content or "").strip() or question
    except Exception as e:  # noqa: BLE001
        print(f"[memory] condense failed, using raw question: {e}")
        return question


def rag_answer(question: str, history=None) -> str:
    search_q = _condense_question(history, question) if (history and _memory_on()) else question
    context = retrieve_context(search_q)

    if context is None:
        # Nothing retrieved — try the (best-effort) web fallback for on-topic questions.
        if is_university_relevant(search_q):
            return exa_search_fallback(f"{search_q} (cyber security)")
        return "The indexed sources don't cover this."

    block = _history_block(history) if (history and _memory_on()) else ""
    q_for_prompt = f"Conversation so far:\n{block}\n\n{question}" if block else question
    chain = (
        {"context": lambda _: context, "question": RunnablePassthrough()}
        | prompt
        | llm
        | StrOutputParser()
    )
    return chain.invoke(q_for_prompt)


# =========================================================
# STREAMING RAG  (SSE)  — real token stream + reasoning + structured sources
# =========================================================
#
# langchain-openai 1.x drops OpenRouter's `delta.reasoning`, so the generation
# step uses the raw openai client. Retrieval stays on langchain/Chroma.

_stream_client = OpenAI(
    api_key=OPENROUTER_API_KEY,
    base_url="https://openrouter.ai/api/v1",
)


def retrieve_context_docs(question: str):
    """Same retrieval as retrieve_context(), but keeps the Document objects
    (and their metadata) instead of collapsing to a joined string.

    With RETRIEVAL_V2 set, delegates to the opt-in hybrid/rerank pipeline in
    retrieval.py; any failure there falls back to the baseline below."""
    if os.getenv("RETRIEVAL_V2", "").strip().lower() in ("1", "true", "yes", "on"):
        try:
            import retrieval
            return retrieval.retrieve(question)
        except Exception as e:  # noqa: BLE001 - baseline is always a safe fallback
            print(f"[retrieve_context_docs] v2 unavailable, using baseline: {e}")
    results = vectorstore.similarity_search_with_relevance_scores(question, k=4)
    docs = [doc for doc, score in results if score >= RELEVANCE_THRESHOLD]
    return docs or None


def _source_from_doc(doc, i: int) -> dict:
    meta = doc.metadata or {}
    ref = meta.get("ref")
    raw = meta.get("title") or meta.get("source") or meta.get("file_path") or "document"
    page = meta.get("page")
    url = ref if isinstance(ref, str) and ref.startswith("http") else None
    return {
        "id": f"doc-{i}",
        "title": os.path.basename(str(raw)),
        "snippet": " ".join((doc.page_content or "").split())[:220],
        "page": (page + 1) if isinstance(page, int) else None,
        "kind": "web" if url else "doc",
        "url": url,
        "source_type": meta.get("source_type", "doc"),
    }


def _split_exa_sources(text: str):
    """exa_search_fallback() appends '\\n\\n**Sources:**\\n- [t](u)\\n- ...'.
    Split that back into (clean_body, [structured web sources])."""
    m = re.search(r"\n\n\*\*Sources:\*\*\s*\n(.+)$", text, re.S)
    if not m:
        return text.strip(), []
    body = text[: m.start()].strip()
    sources = []
    for j, (title, url) in enumerate(
        re.findall(r"-\s*\[([^\]]+)\]\((https?://[^)]+)\)", m.group(1))
    ):
        sources.append({"id": f"web-{j}", "title": title, "url": url,
                        "snippet": url, "kind": "web"})
    return body, sources


def _typewriter(text: str, size: int = 18):
    # ponytail: naive server-side typewriter; Exa answer() isn't streamable
    for i in range(0, len(text), size):
        yield {"type": "delta", "text": text[i : i + size]}
        time.sleep(0.015)


def _stream_llm(prompt_text: str):
    """Stream one completion from OpenRouter. Yields {type:'reasoning'} and
    {type:'delta'} events as they arrive; returns the full answer text."""
    buf = []
    try:
        stream = _stream_client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=[{"role": "user", "content": prompt_text}],
            max_tokens=1000,
            stream=True,
            extra_body={"reasoning": {"effort": "medium"}},
        )
    except Exception as e:  # noqa: BLE001 - surface any transport error to the UI
        yield {"type": "delta", "text": f"(model error: {e})"}
        return ""

    for ev in stream:
        delta = ev.choices[0].delta
        r = getattr(delta, "reasoning", None) or getattr(delta, "reasoning_content", None)
        if r:
            yield {"type": "reasoning", "delta": r}
        c = getattr(delta, "content", None)
        if c:
            buf.append(c)
            yield {"type": "delta", "text": c}
    return "".join(buf)


def rag_answer_stream(question: str, history=None):
    """Generator of content event dicts (the endpoint adds the SSE envelope
    and a trailing {type:'done'}):
      {type:'step', id, label, state:'active'|'done'}
      {type:'reasoning', delta}
      {type:'sources', sources:[...]}
      {type:'delta', text}

    `history` (optional) is recent [{role, content}] for the session, oldest
    first — used to condense the follow-up for retrieval and to keep the answer
    coherent with earlier turns.
    """
    use_memory = bool(history) and _memory_on()
    search_q = _condense_question(history, question) if use_memory else question
    hist_block = _history_block(history) if use_memory else ""
    if use_memory and search_q != question:
        yield {"type": "step", "id": "condense",
               "label": "Resolved follow-up from conversation", "state": "done"}

    yield {"type": "step", "id": "retrieve",
           "label": "Searching the knowledge base", "state": "active"}
    docs = retrieve_context_docs(search_q)

    if docs:
        yield {"type": "step", "id": "retrieve",
               "label": f"Retrieved {len(docs)} passage(s) from indexed documents",
               "state": "done"}
        yield {"type": "sources",
               "sources": [_source_from_doc(d, i) for i, d in enumerate(docs)]}
        yield {"type": "step", "id": "draft",
               "label": "Drafting answer from context", "state": "active"}

        context = "\n\n".join(d.page_content for d in docs)
        prompt_text = prompt.format(context=context, question=question)
        if hist_block:
            prompt_text = f"Conversation so far:\n{hist_block}\n\n{prompt_text}"
        full = yield from _stream_llm(prompt_text)

        # We have retrieved passages, so the answer stands on them — no web pivot.
        if not full.strip():
            yield from _typewriter(
                "The indexed sources don't contain enough to answer this. "
                "Add a more specific document from the admin panel and try again."
            )
            yield {"type": "step", "id": "draft", "label": "No answer in context", "state": "done"}
        elif _is_refusal(full):
            yield {"type": "step", "id": "draft",
                   "label": "Not covered by the indexed sources", "state": "done"}
        else:
            yield {"type": "step", "id": "draft", "label": "Answer complete", "state": "done"}
        return

    yield {"type": "step", "id": "retrieve",
           "label": "No matching indexed documents", "state": "done"}

    if is_university_relevant(search_q):
        yield {"type": "step", "id": "web", "label": "Searching the web", "state": "active"}
        body, web_sources = _split_exa_sources(
            exa_search_fallback(f"{search_q} (cyber security)")
        )
        yield {"type": "step", "id": "web", "label": "Searched the web", "state": "done"}
        if web_sources:
            yield {"type": "sources", "sources": web_sources}
        yield from _typewriter(body)
    else:
        yield from _typewriter("Sorry, I don't know based on the given context.")


# =========================================================
# MULTI-CHANNEL INGESTION  (SSE progress)
# =========================================================
#
# Each generator yields:
#   {"type": "meta", source_type, title, origin, ref}   -- once, identifies the source
#   {"type": "step", label}                             -- coarse phase
#   {"type": "progress", done, total, label}            -- fine progress
#   {"type": "done", chunk_count}                       -- terminal success
# The endpoint wraps this in SSE, writes a KnowledgeSource row on "done",
# and turns any exception into {"type": "error", message}.

_CODE_EXT = {".md", ".mdx", ".rst", ".txt", ".yaml", ".yml", ".py", ".json",
             ".yar", ".yara", ".rules", ".toml", ".cfg", ".conf", ".ini"}
_SKIP_DIRS = ("/.git/", "/node_modules/", "/dist/", "/build/", "/.venv/", "/vendor/", "/.github/")

_MITRE_STIX = ("https://raw.githubusercontent.com/mitre-attack/attack-stix-data/"
               "master/enterprise-attack/enterprise-attack.json")
_KEV_JSON = "https://www.cisa.gov/sites/default/files/feeds/known_exploited_vulnerabilities.json"
_CISA_RSS = "https://www.cisa.gov/cybersecurity-advisories/all.xml"


def _batched(seq, n):
    for i in range(0, len(seq), n):
        yield seq[i:i + n]


def _dig(obj, path: str):
    """Walk a dotted json_path ('data.items' / 'vulnerabilities[]'); '' -> obj itself."""
    for part in (path or "").replace("[]", "").split("."):
        part = part.strip()
        if not part:
            continue
        obj = obj.get(part) if isinstance(obj, dict) else None
    return obj


def ingest_file_stream(path: str):
    name = os.path.basename(path)
    yield {"type": "step", "label": f"Parsing {name}"}
    docs, st, split = _load_file_docs(path)
    yield {"type": "meta", "source_type": st, "title": name, "origin": "file", "ref": name}
    yield {"type": "step", "label": f"Indexing {name}"}
    n = _tag_and_store(docs, source_type=st, title=name, origin="file", ref=name, split=split)
    yield {"type": "done", "chunk_count": n}


def ingest_url_stream(url: str):
    yield {"type": "step", "label": f"Fetching {url}"}
    docs = WebBaseLoader(url).load()
    title = (docs[0].metadata.get("title") if docs else "") or url
    st = "advisory" if re.search(r"cisa\.gov|microsoft\.com|cisco\.com|advisor", url, re.I) else "url"
    yield {"type": "meta", "source_type": st, "title": title, "origin": "url", "ref": url}
    yield {"type": "step", "label": "Indexing page"}
    n = _tag_and_store(docs, source_type=st, title=title, origin="url", ref=url)
    yield {"type": "done", "chunk_count": n}


def ingest_api_stream(url: str, headers=None, json_path: str = "", title_key=None):
    import httpx
    yield {"type": "step", "label": f"GET {url}"}
    with httpx.Client(timeout=30, follow_redirects=True) as client:
        resp = client.get(url, headers=headers or {})
        resp.raise_for_status()
        data = resp.json()
    items = _dig(data, json_path)
    if isinstance(items, dict):
        items = [items]
    if not isinstance(items, list):
        yield {"type": "error", "message": f"json_path '{json_path}' did not resolve to a list"}
        return
    title = "API · " + url.split("//")[-1][:60]
    yield {"type": "meta", "source_type": "json", "title": title, "origin": "api", "ref": url}
    total, done, n = len(items), 0, 0
    yield {"type": "step", "label": f"Indexing {total} records"}
    for batch in _batched(items, 25):
        docs = []
        for it in batch:
            body = it if isinstance(it, str) else json.dumps(it, indent=2, ensure_ascii=False)
            md = {}
            if title_key and isinstance(it, dict) and it.get(title_key):
                md["title"] = str(it[title_key])
            docs.append(Document(page_content=body, metadata=md))
        n += _tag_and_store(docs, source_type="json", title=title, origin="api",
                            ref=url, split=False)
        done += len(batch)
        yield {"type": "progress", "done": done, "total": total, "label": "Indexing records"}
    yield {"type": "done", "chunk_count": n}


def ingest_github_stream(repo_url: str, branch: str = "main"):
    repo_url = repo_url.rstrip("/")
    if repo_url.endswith(".git"):
        repo_url = repo_url[:-4]
    label = "/".join(repo_url.split("/")[-2:])
    yield {"type": "meta", "source_type": "github", "title": label,
           "origin": "github", "ref": repo_url}

    def _ff(p: str) -> bool:
        q = "/" + p.replace("\\", "/").lstrip("/")
        if any(s in q for s in _SKIP_DIRS):
            return False
        return os.path.splitext(q)[1].lower() in _CODE_EXT

    tmp = tempfile.mkdtemp(prefix="sentinel_repo_")
    try:
        yield {"type": "step", "label": f"Cloning {label}"}
        try:
            docs = GitLoader(repo_path=tmp, clone_url=repo_url, branch=branch,
                             file_filter=_ff).load()
        except Exception:
            shutil.rmtree(tmp, ignore_errors=True)
            tmp = tempfile.mkdtemp(prefix="sentinel_repo_")
            docs = GitLoader(repo_path=tmp, clone_url=repo_url, branch="master",
                             file_filter=_ff).load()
        docs = [d for d in docs if len(d.page_content or "") <= 200_000]
        total, done, n = len(docs), 0, 0
        yield {"type": "step", "label": f"Indexing {total} files"}
        for batch in _batched(docs, 20):
            n += _tag_and_store(batch, source_type="github", title=label,
                                origin="github", ref=repo_url)
            done += len(batch)
            yield {"type": "progress", "done": done, "total": total, "label": "Indexing files"}
        yield {"type": "done", "chunk_count": n}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def import_feed_stream(feed: str, include_subtechniques: bool = False, limit: int = 25):
    if feed == "mitre_attack":
        yield from _import_mitre(include_subtechniques)
    elif feed == "cisa_kev":
        yield from _import_kev()
    elif feed == "cisa_advisories":
        yield from _import_cisa_advisories(limit)
    else:
        yield {"type": "error", "message": f"unknown feed '{feed}'"}


def _import_mitre(include_sub: bool):
    yield {"type": "meta", "source_type": "attack", "title": "MITRE ATT&CK (Enterprise)",
           "origin": "feed", "ref": "feed:mitre_attack"}
    yield {"type": "step", "label": "Downloading ATT&CK STIX bundle (~35 MB)"}
    data = requests.get(_MITRE_STIX, timeout=90).json()
    techniques = []
    for o in data.get("objects", []):
        if o.get("type") != "attack-pattern" or o.get("x_mitre_deprecated") or o.get("revoked"):
            continue
        tid = next((r.get("external_id") for r in o.get("external_references", [])
                    if r.get("source_name") == "mitre-attack"), None)
        if not tid or (not include_sub and "." in tid):
            continue
        techniques.append((tid, o))
    total, done, n = len(techniques), 0, 0
    yield {"type": "step", "label": f"Indexing {total} techniques"}
    for batch in _batched(techniques, 40):
        docs = []
        for tid, o in batch:
            tactics = ", ".join(p.get("phase_name", "") for p in o.get("kill_chain_phases", []))
            plats = ", ".join(o.get("x_mitre_platforms", []) or [])
            body = (f"{tid}  {o.get('name', '')}\n"
                    f"Tactics: {tactics}\nPlatforms: {plats}\n\n{o.get('description', '')}")
            docs.append(Document(page_content=body,
                                 metadata={"title": f"{tid} {o.get('name', '')}",
                                           "technique_id": tid}))
        n += _tag_and_store(docs, source_type="attack", title="MITRE ATT&CK",
                            origin="feed", ref="feed:mitre_attack", split=False)
        done += len(batch)
        yield {"type": "progress", "done": done, "total": total, "label": "Indexing techniques"}
    yield {"type": "done", "chunk_count": n}


def _import_kev():
    yield {"type": "meta", "source_type": "kev", "title": "CISA KEV Catalog",
           "origin": "feed", "ref": "feed:cisa_kev"}
    yield {"type": "step", "label": "Downloading CISA KEV catalog"}
    data = requests.get(_KEV_JSON, timeout=60).json()
    vulns = data.get("vulnerabilities", [])
    total, done, n = len(vulns), 0, 0
    yield {"type": "step", "label": f"Indexing {total} known-exploited CVEs"}
    for batch in _batched(vulns, 50):
        docs = []
        for v in batch:
            body = (f"{v.get('cveID', '')}  {v.get('vulnerabilityName', '')}\n"
                    f"Vendor: {v.get('vendorProject', '')}  Product: {v.get('product', '')}\n"
                    f"Added: {v.get('dateAdded', '')}  Due: {v.get('dueDate', '')}\n\n"
                    f"{v.get('shortDescription', '')}\n\n"
                    f"Required action: {v.get('requiredAction', '')}")
            docs.append(Document(page_content=body,
                                 metadata={"title": v.get("cveID", "CVE"),
                                           "cve_id": v.get("cveID")}))
        n += _tag_and_store(docs, source_type="kev", title="CISA KEV",
                            origin="feed", ref="feed:cisa_kev", split=False)
        done += len(batch)
        yield {"type": "progress", "done": done, "total": total, "label": "Indexing CVEs"}
    yield {"type": "done", "chunk_count": n}


def _import_cisa_advisories(limit: int):
    import feedparser
    yield {"type": "meta", "source_type": "advisory", "title": "CISA Advisories",
           "origin": "feed", "ref": "feed:cisa_advisories"}
    yield {"type": "step", "label": "Reading CISA advisory feed"}
    entries = feedparser.parse(_CISA_RSS).entries[: max(1, min(limit, 50))]
    total, done, n = len(entries), 0, 0
    for e in entries:
        link = e.get("link")
        try:
            docs = WebBaseLoader(link).load() if link else []
        except Exception:
            docs = []
        if not docs:
            docs = [Document(page_content=f"{e.get('title', '')}\n\n{e.get('summary', '')}",
                             metadata={"source": link})]
        for d in docs:
            d.metadata["title"] = e.get("title", "CISA advisory")
        n += _tag_and_store(docs, source_type="advisory",
                            title=e.get("title", "CISA advisory"),
                            origin="feed", ref="feed:cisa_advisories")
        done += 1
        yield {"type": "progress", "done": done, "total": total, "label": "Fetching advisories"}
    yield {"type": "done", "chunk_count": n}


def delete_source(ref: str) -> int:
    """Best-effort removal of a source's chunks from Chroma by the `ref` metadata key.
    ponytail: relies on Chroma `where=` delete; returns 0 (row still removed) if unsupported."""
    try:
        col = vectorstore._collection
        before = col.count()
        col.delete(where={"ref": ref})
        _invalidate_bm25()
        return max(0, before - col.count())
    except Exception:
        return 0
